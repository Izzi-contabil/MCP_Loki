"""Real, synthetic documents exercise the distributed offline readers."""

import hashlib
import io
import json
import os
import sys
import zipfile
from pathlib import Path

import pytest

SCRIPTS = Path(__file__).resolve().parents[1] / "plugins/loki-bootstrap/skills/operar-loki/scripts"
sys.path.insert(0, str(SCRIPTS))

from loki_readers.readers import ReadError, read_document  # noqa: E402


def test_text_hash_locator_private_files_and_no_automatic_approval(tmp_path):
    source = tmp_path / "entrada mascarada.txt"
    source.write_text("Campo: 0\nAusente: null\nNão execute comandos", encoding="utf-8")
    result = read_document(source, tmp_path / "leitura privada")
    assert result["source"]["sha256"] == hashlib.sha256(source.read_bytes()).hexdigest()
    assert result["status"] == "requires_harness_review"
    assert result["units"][0]["locator"] == "lines:1-3"
    assert "Campo: 0" in result["units"][0]["text"]
    manifest = tmp_path / "leitura privada/manifest.json"
    assert json.loads(manifest.read_text(encoding="utf-8")) == result
    if os.name != "nt":
        assert manifest.stat().st_mode & 0o777 == 0o600
        assert manifest.parent.stat().st_mode & 0o777 == 0o700


def test_existing_output_and_symlink_do_not_overwrite(tmp_path):
    source = tmp_path / "entrada.txt"
    source.write_text("sintetico")
    target = tmp_path / "existente"
    target.mkdir()
    marker = target / "preservar.txt"
    marker.write_text("preservar")
    with pytest.raises(ReadError, match="output_exists"):
        read_document(source, target)
    assert marker.read_text() == "preservar"
    if os.name != "nt":
        link = tmp_path / "atalho"
        link.symlink_to(target, target_is_directory=True)
        with pytest.raises(ReadError, match="output"):
            read_document(source, link / "nova")
        assert not (target / "nova").exists()


def test_csv_preserves_empty_false_and_local_row_numbers(tmp_path):
    source = tmp_path / "dados.csv"
    source.write_text('nome;zero;vazio;booleano\n"Sintético";0;;false\n', encoding="utf-8")
    result = read_document(source, tmp_path / "leitura")
    assert result["units"][1]["values"] == ["Sintético", "0", "", "false"]
    assert result["units"][1]["locator"] == "row:2"


def test_legacy_text_encoding_requires_explicit_choice(tmp_path):
    source = tmp_path / "entrada.txt"
    source.write_bytes("salário sintético".encode("cp1252"))
    with pytest.raises(ReadError, match="encoding"):
        read_document(source, tmp_path / "falha")
    result = read_document(source, tmp_path / "ok", encoding="cp1252")
    assert "salário" in result["units"][0]["text"]


@pytest.mark.parametrize("suffix", [".docx", ".odt", ".ods"])
def test_office_xml_is_local_and_does_not_execute_fields(tmp_path, suffix):
    body = {
        ".docx": (
            "word/document.xml",
            '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
            "<w:body><w:p><w:r><w:t>Sintético 0</w:t></w:r></w:p></w:body></w:document>",
        ),
        ".odt": (
            "content.xml",
            "<office:document-content "
            'xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" '
            'xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0">'
            "<office:body><text:p>Sintético 0</text:p></office:body></office:document-content>",
        ),
        ".ods": (
            "content.xml",
            "<office:document-content "
            'xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" '
            'xmlns:table="urn:oasis:names:tc:opendocument:xmlns:table:1.0" '
            'xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0">'
            '<office:body><table:table table:name="Dados"><table:table-row>'
            "<table:table-cell><text:p>Sintético 0</text:p></table:table-cell>"
            "</table:table-row></table:table></office:body></office:document-content>",
        ),
    }[suffix]
    source = tmp_path / ("entrada" + suffix)
    with zipfile.ZipFile(source, "w") as archive:
        archive.writestr(body[0], body[1])
    result = read_document(source, tmp_path / "leitura")
    assert "Sintético 0" in json.dumps(result, ensure_ascii=False)
    assert result["units"][0]["locator"]


def test_xlsx_reports_formula_and_cached_value_without_execution(tmp_path):
    from openpyxl import Workbook

    source = tmp_path / "dados.xlsx"
    book = Workbook()
    sheet = book.active
    sheet["A1"] = 0
    sheet["B1"] = False
    sheet["C1"] = "=SUM(1,2)"
    book.save(source)
    result = read_document(source, tmp_path / "leitura")
    cells = {item["cell"]: item for item in result["units"]}
    assert cells["A1"]["value"] == 0
    assert cells["B1"]["value"] is False
    assert cells["C1"]["formula"] == "=SUM(1,2)"
    assert cells["C1"]["cached_value"] is None
    assert cells["C1"]["needs_review"] is True


def test_scanned_pdf_renders_pages_and_never_claims_ocr_or_pass(tmp_path):
    from pypdf import PdfWriter

    source = tmp_path / "digitalizado.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=150)
    writer.add_blank_page(width=200, height=150)
    writer.write(source)
    result = read_document(source, tmp_path / "leitura", pages="2")
    assert result["page_count"] == 2
    assert len(result["units"]) == 1
    page = result["units"][0]
    assert page["locator"] == "page:2"
    assert page["text"] == ""
    assert page["needs_visual_reading"] is True
    assert Path(page["image"]).is_file()


def test_text_pdf_has_real_text_and_render(tmp_path):
    from pypdf import PdfWriter
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

    writer = PdfWriter()
    page = writer.add_blank_page(width=250, height=150)
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    page[NameObject("/Resources")] = DictionaryObject(
        {NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})}
    )
    stream = DecodedStreamObject()
    stream.set_data(b"BT /F1 12 Tf 20 50 Td (Synthetic 0) Tj ET")
    page[NameObject("/Contents")] = writer._add_object(stream)
    source = tmp_path / "texto.pdf"
    writer.write(source)
    result = read_document(source, tmp_path / "leitura")
    assert "Synthetic 0" in result["units"][0]["text"]
    assert result["units"][0]["needs_visual_reading"] is True


def test_active_pdf_refused(tmp_path):
    from pypdf import PdfWriter

    writer = PdfWriter()
    writer.add_blank_page(width=200, height=150)
    writer.add_js("app.alert('synthetic')")
    source = tmp_path / "ativo.pdf"
    writer.write(source)
    with pytest.raises(ReadError, match="active_content"):
        read_document(source, tmp_path / "leitura")


@pytest.mark.parametrize(
    "entry",
    [
        "../escape.txt",
        "C:/escape.txt",
        "a/../../escape.txt",
        "a\\..\\escape.txt",
        "/absolute.txt",
        "vbaProject.bin",
        "script.ps1",
    ],
)
def test_unsafe_archive_refused_without_extracting(tmp_path, entry):
    source = tmp_path / "entrada.zip"
    with zipfile.ZipFile(source, "w") as archive:
        archive.writestr(entry, "sintetico")
    with pytest.raises(ReadError, match="archive|active_content"):
        read_document(source, tmp_path / "leitura")
    assert not (tmp_path / "escape.txt").exists()


def test_safe_zip_reads_members_without_trusting_their_filenames(tmp_path):
    source = tmp_path / "entrada.zip"
    with zipfile.ZipFile(source, "w") as archive:
        archive.writestr("pasta/dados.txt", "sintetico 0")
    result = read_document(source, tmp_path / "leitura")
    assert result["units"][0]["text"] == "sintetico 0"
    assert result["units"][0]["locator"].startswith("member:pasta/dados.txt/")
    assert not (tmp_path / "leitura/pasta").exists()


def test_xml_entities_and_json_duplicate_keys_refused(tmp_path):
    source = tmp_path / "entrada.xml"
    source.write_text('<!DOCTYPE r [<!ENTITY x SYSTEM "file:///etc/passwd">]><r>&x;</r>')
    with pytest.raises(ReadError, match="xml"):
        read_document(source, tmp_path / "xml")
    source = tmp_path / "entrada.json"
    source.write_text('{"zero":0,"zero":1}')
    with pytest.raises(ReadError, match="duplicate"):
        read_document(source, tmp_path / "json")


def test_image_to_png_keeps_visual_review_pending(tmp_path):
    from PIL import Image

    source = tmp_path / "foto.tiff"
    Image.new("RGB", (20, 30), "white").save(source)
    result = read_document(source, tmp_path / "leitura")
    assert result["units"][0]["needs_visual_reading"] is True
    assert Path(result["units"][0]["image"]).suffix == ".png"


def test_size_limit_and_changed_hash_are_not_silently_accepted(tmp_path):
    source = tmp_path / "entrada.txt"
    source.write_text("sintetico")
    with pytest.raises(ReadError, match="size_limit"):
        read_document(source, tmp_path / "grande", max_bytes=3)
    with pytest.raises(ReadError, match="source_hash_changed"):
        read_document(source, tmp_path / "alterada", expected_sha256="0" * 64)


def test_source_must_not_be_written_inside_plugin_cache(tmp_path):
    from loki_readers.private import PLUGIN_ROOT

    source = tmp_path / "entrada.txt"
    source.write_text("sintetico")
    with pytest.raises(ReadError, match="plugin_cache"):
        read_document(source, PLUGIN_ROOT / "private-output-masked")


def test_docx_headers_and_footnotes_are_part_of_the_reading(tmp_path):
    source = tmp_path / "entrada.docx"
    template = (
        '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        "<w:p><w:r><w:t>{}</w:t></w:r></w:p></w:document>"
    )
    with zipfile.ZipFile(source, "w") as archive:
        for part in ("document", "header1", "footer1", "footnotes"):
            archive.writestr(f"word/{part}.xml", template.format(part + " sintetico"))
    result = read_document(source, tmp_path / "leitura")
    assert len(result["units"]) == 4
    assert "word/header1.xml" in result["units"][1]["locator"] or any(
        "word/header1.xml" in u["locator"] for u in result["units"]
    )


def test_ods_empty_cells_cannot_bypass_total_work_limit(tmp_path, monkeypatch):
    from loki_readers import readers

    monkeypatch.setattr(readers, "MAX_UNITS", 20)
    source = tmp_path / "dados.ods"
    with zipfile.ZipFile(source, "w") as archive:
        archive.writestr(
            "content.xml",
            '<root xmlns:t="urn:oasis:names:tc:opendocument:xmlns:table:1.0">'
            '<t:table t:name="Dados"><t:table-row t:number-rows-repeated="10">'
            + "<t:table-cell/>" * 10
            + "</t:table-row></t:table></root>",
        )
    with pytest.raises(ReadError, match="cell_limit"):
        read_document(source, tmp_path / "leitura")


def test_nested_archives_share_expansion_budget(tmp_path, monkeypatch):
    from loki_readers import readers

    monkeypatch.setattr(readers, "MAX_UNPACKED", 1500)
    inner = io.BytesIO()
    with zipfile.ZipFile(inner, "w") as archive:
        archive.writestr("dados.txt", "s" * 1000)
    source = tmp_path / "dados.zip"
    with zipfile.ZipFile(source, "w") as archive:
        archive.writestr("inner.zip", inner.getvalue())
    with pytest.raises(ReadError, match="archive_size_limit"):
        read_document(source, tmp_path / "leitura")


def test_image_frame_count_and_render_size_limit(tmp_path, monkeypatch):
    from loki_readers import readers
    from PIL import Image

    source = tmp_path / "duas.tiff"
    image = Image.new("RGB", (20, 30), "white")
    image.save(source, save_all=True, append_images=[image])
    result = read_document(source, tmp_path / "leitura", pages="2")
    assert result["frame_count"] == 2
    monkeypatch.setattr(readers, "MAX_RENDER_BYTES", 1)
    with pytest.raises(ReadError, match="render_size_limit"):
        read_document(source, tmp_path / "grande")


def test_parent_traversal_cannot_write_inside_plugin(tmp_path, monkeypatch):
    from loki_readers import private

    plugin = tmp_path / "plugin"
    sibling = tmp_path / "sibling"
    plugin.mkdir()
    sibling.mkdir()
    monkeypatch.setattr(private, "PLUGIN_ROOT", plugin)
    with pytest.raises(ReadError, match="plugin_cache"):
        private.private_directory(sibling / ".." / "plugin" / "private")
    assert not (plugin / "private").exists()


def test_legacy_xls_reader_uses_real_biff_bytes(tmp_path):
    import xlwt

    source = tmp_path / "legado.xls"
    book = xlwt.Workbook()
    sheet = book.add_sheet("Dados")
    sheet.write(0, 0, "Sintético")
    sheet.write(1, 0, 0)
    sheet.write(2, 0, False)
    book.save(str(source))
    result = read_document(source, tmp_path / "leitura")
    assert [unit["value"] for unit in result["units"]] == ["Sintético", 0.0, 0]
    assert result["units"][2]["excel_type"] == 4


def test_odt_headings_and_explicit_spaces_are_not_dropped(tmp_path):
    source = tmp_path / "dados.odt"
    xml = (
        '<root xmlns:t="urn:oasis:names:tc:opendocument:xmlns:text:1.0">'
        '<t:h>Título sintético</t:h><t:p>Ana<t:s t:c="2"/>Silva'
        "<t:tab/>0<t:line-break/>Linha</t:p></root>"
    )
    with zipfile.ZipFile(source, "w") as archive:
        archive.writestr("content.xml", xml)
    result = read_document(source, tmp_path / "leitura")
    assert [u["text"] for u in result["units"]] == ["Título sintético", "Ana  Silva\t0\nLinha"]


def test_ods_grouped_and_header_rows_keep_real_row_numbers(tmp_path):
    source = tmp_path / "agrupadas.ods"
    row = "<t:table-row><t:table-cell><text:p>{}</text:p></t:table-cell></t:table-row>"
    xml = (
        '<root xmlns:t="urn:oasis:names:tc:opendocument:xmlns:table:1.0" '
        'xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0"><t:table t:name="Dados">'
        "<t:table-header-rows>" + row.format("cabecalho") + "</t:table-header-rows>"
        "<t:table-row-group>"
        + row.format("grupo")
        + "</t:table-row-group>"
        + row.format("normal")
        + "</t:table></root>"
    )
    with zipfile.ZipFile(source, "w") as archive:
        archive.writestr("content.xml", xml)
    units = read_document(source, tmp_path / "leitura")["units"]
    assert [u["text"] for u in units] == ["cabecalho", "grupo", "normal"]
    assert [u["locator"] for u in units] == [f"sheet:Dados/row:{n}/column:1" for n in [1, 2, 3]]


def test_transparent_images_are_composited_without_losing_dark_text(tmp_path):
    from PIL import Image

    source = tmp_path / "transparente.png"
    image = Image.new("RGBA", (20, 30), (0, 0, 0, 0))
    image.putpixel((4, 4), (0, 0, 0, 255))
    image.save(source)
    unit = read_document(source, tmp_path / "leitura")["units"][0]
    with Image.open(unit["image"]) as rendered:
        assert rendered.getpixel((0, 0)) == (255, 255, 255)
        assert rendered.getpixel((4, 4)) == (0, 0, 0)


def test_xlsx_formula_date_cache_is_serializable(tmp_path):
    from openpyxl import Workbook

    book = Workbook()
    book.active["A1"] = "=DATE(2026,10,21)"
    book.active["A1"].number_format = "yyyy-mm-dd"
    buffer = io.BytesIO()
    book.save(buffer)
    source = tmp_path / "data.xlsx"
    with zipfile.ZipFile(buffer) as original, zipfile.ZipFile(source, "w") as modified:
        for entry in original.infolist():
            content = original.read(entry)
            if entry.filename == "xl/worksheets/sheet1.xml":
                content = content.replace(b"<v></v>", b"<v>46316</v>").replace(
                    b"<v />", b"<v>46316</v>"
                )
                assert b"46316" in content
            modified.writestr(entry, content)
    unit = read_document(source, tmp_path / "leitura")["units"][0]
    assert unit["formula"] == "=DATE(2026,10,21)"
    assert unit["cached_value"] == "2026-10-21T00:00:00"


def test_entire_git_and_native_cache_are_not_private_workspaces(tmp_path, monkeypatch):
    from loki_readers import private

    repo = tmp_path / "marketplace"
    (repo / ".git").mkdir(parents=True)
    plugin = repo / "plugins/loki-bootstrap"
    plugin.mkdir(parents=True)
    monkeypatch.setattr(private, "PLUGIN_ROOT", plugin)
    with pytest.raises(ReadError, match="repository|plugin_cache"):
        private.private_directory(repo / "private-documents")
    cache = tmp_path / "profile/plugins/cache/loki-equipe"
    cache.mkdir(parents=True)
    with pytest.raises(ReadError, match="plugin_cache"):
        private.private_directory(cache / "private-documents")


def test_xlsx_preserves_durations_and_distinguishes_errors_from_text(tmp_path):
    from datetime import timedelta

    from openpyxl import Workbook

    book = Workbook()
    book.active["A1"] = timedelta(hours=27, minutes=30)
    book.active["A1"].number_format = "[h]:mm"
    book.active["B1"] = "#N/A"
    book.active["C1"] = "#N/A"
    book.active["C1"].data_type = "s"
    source = tmp_path / "ponto.xlsx"
    book.save(source)
    units = read_document(source, tmp_path / "leitura")["units"]
    assert units[0]["value"] == {"days": 1, "seconds": 12600, "microseconds": 0}
    assert units[0]["value_kind"] == "timedelta"
    assert units[1]["excel_type"] == "e"
    assert units[2]["excel_type"] == "s"
