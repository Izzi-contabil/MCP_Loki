"""Bounded local text, table and image readers. No networking, OCR service or LLM."""

import csv
import hashlib
import io
import re
import stat
import warnings
import zipfile
from pathlib import Path, PurePosixPath, PureWindowsPath
from typing import Any

from .private import (
    ReadError,
    checked_bytes,
    private_directory,
    strict_json,
    write_bytes,
    write_json,
)

MAX_BYTES = 32 * 1024 * 1024
MAX_UNPACKED = 64 * 1024 * 1024
MAX_UNITS = 200000
MAX_RENDER_BYTES = 128 * 1024 * 1024


def safe_archive(content: bytes, budget: dict[str, int] | None = None) -> zipfile.ZipFile:
    archive = zipfile.ZipFile(io.BytesIO(content))
    entries = archive.infolist()
    budget = budget if budget is not None else {"unpacked": 0}
    budget["unpacked"] += sum(e.file_size for e in entries)
    if not entries or len(entries) > 10000 or budget["unpacked"] > MAX_UNPACKED:
        raise ReadError("archive_size_limit")
    names: set[str] = set()
    for entry in entries:
        name = entry.filename
        path = PurePosixPath(name)
        if (
            path.is_absolute()
            or PureWindowsPath(name).drive
            or ".." in path.parts
            or "\\" in name
            or name in names
            or entry.flag_bits & 1
            or stat.S_ISLNK(entry.external_attr >> 16)
            or entry.file_size > MAX_UNPACKED
            or entry.file_size > 500 * max(entry.compress_size, 1)
        ):
            raise ReadError("unsafe_archive")
        names.add(name)
        lower = name.lower()
        if any(
            token in lower for token in ("vbaproject", "activex", "embeddings/", "scripts/")
        ) or lower.endswith((".exe", ".dll", ".ps1", ".js", ".bat", ".vbs", ".cmd")):
            raise ReadError("active_content")
        if lower.endswith((".xml", ".rels")):
            from defusedxml.ElementTree import fromstring

            try:
                fromstring(archive.read(entry), forbid_dtd=True)
            except Exception:
                raise ReadError("unsafe_xml") from None
    return archive


def _preview(path: Path, content: bytes) -> None:
    previews = list(path.parent.glob("*.png"))
    if (
        len(previews) >= 200
        or len(content) + sum(p.stat().st_size for p in previews) > MAX_RENDER_BYTES
    ):
        raise ReadError("render_size_limit")
    write_bytes(path, content)


def _selected_pages(spec: str | None, count: int) -> list[int]:
    if count > 2000 or count < 1:
        raise ReadError("page_limit")
    selected: set[int] = set()
    if spec:
        for item in spec.split(","):
            if not re.fullmatch(r"[1-9][0-9]*(?:-[1-9][0-9]*)?", item):
                raise ReadError("invalid_pages")
            bounds = [int(n) for n in item.split("-")]
            low, high = bounds[0], bounds[-1]
            if not 1 <= low <= high <= count:
                raise ReadError("invalid_pages")
            selected.update(range(low - 1, high))
    else:
        selected.update(range(count))
    if len(selected) > 200:
        raise ReadError("page_batch_limit_use_pages")
    return sorted(selected)


def _pdf(
    content: bytes, output: Path, pages: str | None, prefix: str
) -> tuple[list[dict[str, Any]], int]:
    import pypdfium2 as pdfium
    from pypdf import PdfReader
    from pypdf.generic import DictionaryObject, IndirectObject

    reader = PdfReader(io.BytesIO(content), strict=True)
    if reader.is_encrypted:
        raise ReadError("encrypted_document")
    pending: list[Any] = [reader.trailer]
    seen: set[int] = set()
    for _ in range(100000):
        if not pending:
            break
        node = pending.pop()
        if isinstance(node, IndirectObject):
            node = node.get_object()
        if id(node) in seen:
            continue
        seen.add(id(node))
        if isinstance(node, DictionaryObject):
            if set(node) & {
                "/JS",
                "/JavaScript",
                "/Launch",
                "/EmbeddedFiles",
                "/AA",
                "/RichMedia",
                "/XFA",
            }:
                raise ReadError("active_content")
            if node.get("/S") in {"/JavaScript", "/Launch", "/SubmitForm", "/ImportData"}:
                raise ReadError("active_content")
            pending.extend(node.values())
        elif isinstance(node, list):
            pending.extend(node)
    else:
        raise ReadError("pdf_structure_limit")
    count = len(reader.pages)
    selected = _selected_pages(pages, count)
    units = []
    with pdfium.PdfDocument(content) as document:
        for index in selected:
            page = document[index]
            try:
                width, height = page.get_size()
                if width <= 0 or height <= 0:
                    raise ReadError("invalid_page_size")
                bitmap = page.render(scale=min(2.0, 4000 / max(width, height)))
                try:
                    image = bitmap.to_pil()
                    buffer = io.BytesIO()
                    image.save(buffer, format="PNG")
                finally:
                    bitmap.close()
                path = output / f"{prefix}page-{index + 1:04d}.png"
                _preview(path, buffer.getvalue())
            finally:
                page.close()
            text = reader.pages[index].extract_text() or ""
            if len(text) > 2_000_000:
                raise ReadError("text_limit")
            units.append(
                {
                    "locator": f"page:{index + 1}",
                    "text": text,
                    "image": str(path),
                    "needs_visual_reading": True,
                }
            )
    return units, count


def _images(
    content: bytes, output: Path, prefix: str, pages: str | None
) -> tuple[list[dict[str, Any]], int]:
    from PIL import Image, ImageOps

    units = []
    with warnings.catch_warnings():
        warnings.simplefilter("error", Image.DecompressionBombWarning)
        with Image.open(io.BytesIO(content)) as image:
            count = getattr(image, "n_frames", 1)
            for index in _selected_pages(pages, count):
                image.seek(index)
                if image.width * image.height > 40_000_000:
                    raise ReadError("image_size_limit")
                frame = ImageOps.exif_transpose(image).convert("RGB")
                buffer = io.BytesIO()
                frame.save(buffer, format="PNG")
                path = output / f"{prefix}image-{index + 1:04d}.png"
                _preview(path, buffer.getvalue())
                units.append(
                    {
                        "locator": f"frame:{index + 1}",
                        "image": str(path),
                        "needs_visual_reading": True,
                    }
                )
    return units, count


def _xlsx(content: bytes, budget: dict[str, int]) -> list[dict[str, Any]]:
    from openpyxl import load_workbook

    safe_archive(content, budget).close()
    book = load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
    cache = load_workbook(io.BytesIO(content), read_only=True, data_only=True, keep_links=False)
    units = []
    visited = 0
    try:
        for sheet in book:
            if sheet.max_row and sheet.max_column and sheet.max_row * sheet.max_column > MAX_UNITS:
                raise ReadError("cell_limit")
            cached = cache[sheet.title]
            for row, cached_row in zip(sheet.iter_rows(), cached.iter_rows(), strict=True):
                for cell, cached_cell in zip(row, cached_row, strict=True):
                    visited += 1
                    if visited > MAX_UNITS:
                        raise ReadError("cell_limit")
                    if cell.value is None:
                        continue
                    value = cell.value
                    unit = {
                        "locator": f"sheet:{sheet.title}!{cell.coordinate}",
                        "sheet": sheet.title,
                        "cell": cell.coordinate,
                        "value": value
                        if isinstance(value, (str, int, float, bool))
                        else value.isoformat(),
                    }
                    if cell.data_type == "f":
                        unit.update(
                            formula=value, cached_value=cached_cell.value, needs_review=True
                        )
                    units.append(unit)
                    if len(units) > MAX_UNITS:
                        raise ReadError("cell_limit")
    finally:
        book.close()
        cache.close()
    return units


def _document_text(element: Any, depth: int = 0) -> str:
    if depth > 100:
        raise ReadError("xml_depth_limit")
    parts = [element.text or ""]
    for child in element:
        kind = child.tag.rsplit("}", 1)[-1]
        if kind == "tab":
            parts.append("\t")
        elif kind in {"line-break", "br", "cr"}:
            parts.append("\n")
        elif child.tag == "{urn:oasis:names:tc:opendocument:xmlns:text:1.0}s":
            count = int(child.get("{urn:oasis:names:tc:opendocument:xmlns:text:1.0}c", "1"))
            if not 1 <= count <= 10000:
                raise ReadError("text_space_limit")
            parts.append(" " * count)
        else:
            parts.append(_document_text(child, depth + 1))
        parts.append(child.tail or "")
    return "".join(parts)


def _office(content: bytes, suffix: str, budget: dict[str, int]) -> list[dict[str, Any]]:
    from defusedxml.ElementTree import fromstring

    with safe_archive(content, budget) as archive:
        if suffix == ".docx":
            units = []
            paragraph = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}p"
            for part in sorted(archive.namelist()):
                if not re.fullmatch(
                    r"word/(?:document|header[0-9]+|footer[0-9]+|footnotes|endnotes|comments)\.xml",
                    part,
                ):
                    continue
                root = fromstring(archive.read(part), forbid_dtd=True)
                for i, element in enumerate(root.iter(paragraph), 1):
                    units.append(
                        {
                            "locator": f"part:{part}/paragraph:{i}",
                            "text": _document_text(element),
                        }
                    )
                    if len(units) > MAX_UNITS:
                        raise ReadError("unit_limit")
            if not units:
                raise ReadError("office_document_empty")
            return units
        member = "word/document.xml" if suffix == ".docx" else "content.xml"
        root = fromstring(archive.read(member), forbid_dtd=True)
    if suffix == ".ods":
        table_ns = "{urn:oasis:names:tc:opendocument:xmlns:table:1.0}"
        result = []
        visited = 0
        for table in root.iter(table_ns + "table"):
            name = table.get(table_ns + "name", "")
            row_number = 0
            for row in table.findall(table_ns + "table-row"):
                repeat = int(row.get(table_ns + "number-rows-repeated", "1"))
                if repeat > MAX_UNITS:
                    raise ReadError("cell_limit")
                for _ in range(repeat):
                    row_number += 1
                    column = 0
                    for cell in row:
                        copies = int(cell.get(table_ns + "number-columns-repeated", "1"))
                        if copies < 1 or column + copies > MAX_UNITS:
                            raise ReadError("cell_limit")
                        text = _document_text(cell)
                        for _ in range(copies):
                            visited += 1
                            if visited > MAX_UNITS:
                                raise ReadError("cell_limit")
                            column += 1
                            if text.strip() or cell.attrib:
                                result.append(
                                    {
                                        "locator": f"sheet:{name}/row:{row_number}/column:{column}",
                                        "text": text,
                                        "formula": cell.get(table_ns + "formula"),
                                        "attributes": dict(cell.attrib),
                                    }
                                )
                        if len(result) > MAX_UNITS:
                            raise ReadError("cell_limit")
        return result
    paragraph_tags = {
        "{urn:oasis:names:tc:opendocument:xmlns:text:1.0}" + kind for kind in ("p", "h")
    }
    return [
        {"locator": f"paragraph:{i}", "text": _document_text(element)}
        for i, element in enumerate((node for node in root.iter() if node.tag in paragraph_tags), 1)
    ]


def _read(
    content: bytes,
    suffix: str,
    output: Path,
    encoding: str,
    pages: str | None,
    prefix: str = "",
    depth: int = 0,
    budget: dict[str, int] | None = None,
) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    budget = budget if budget is not None else {"unpacked": 0}
    if content.startswith(b"%PDF-"):
        units, count = _pdf(content, output, pages, prefix)
        return units, {"page_count": count}
    if suffix in {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".webp"}:
        units, count = _images(content, output, prefix, pages)
        return units, {"frame_count": count}
    if suffix == ".xlsx":
        return _xlsx(content, budget), {}
    if suffix == ".xls":
        import xlrd

        book = xlrd.open_workbook(file_contents=content, on_demand=True)
        units = []
        visited = 0
        try:
            for sheet in book.sheets():
                if sheet.nrows * sheet.ncols > MAX_UNITS:
                    raise ReadError("cell_limit")
                for row in range(sheet.nrows):
                    for column in range(sheet.ncols):
                        visited += 1
                        if visited > MAX_UNITS:
                            raise ReadError("cell_limit")
                        cell = sheet.cell(row, column)
                        if cell.ctype:
                            units.append(
                                {
                                    "locator": (
                                        f"sheet:{sheet.name}/row:{row + 1}/column:{column + 1}"
                                    ),
                                    "value": cell.value,
                                    "excel_type": cell.ctype,
                                    "excel_date_mode": book.datemode,
                                }
                            )
        finally:
            book.release_resources()
        return units, {}
    if suffix in {".docx", ".odt", ".ods"}:
        return _office(content, suffix, budget), {}
    if suffix == ".zip":
        if depth >= 3:
            raise ReadError("archive_depth_limit")
        units = []
        members = []
        with safe_archive(content, budget) as archive:
            for i, entry in enumerate(archive.infolist()):
                if entry.is_dir():
                    continue
                member = archive.read(entry)
                children, metadata = _read(
                    member,
                    PurePosixPath(entry.filename).suffix.lower(),
                    output,
                    encoding,
                    pages,
                    f"{prefix}member-{i:04d}-",
                    depth + 1,
                    budget,
                )
                members.append(
                    {
                        "name": entry.filename,
                        "sha256": hashlib.sha256(member).hexdigest(),
                        **metadata,
                    }
                )
                for child in children:
                    child["locator"] = f"member:{entry.filename}/{child['locator']}"
                    child["member_sha256"] = hashlib.sha256(member).hexdigest()
                    units.append(child)
                    if len(units) > MAX_UNITS:
                        raise ReadError("unit_limit")
        return units, {"members": members}
    if suffix not in {".txt", ".csv", ".md", ".json", ".xml"}:
        raise ReadError("unsupported_format")
    try:
        text = content.decode(encoding, errors="strict")
    except (UnicodeError, LookupError):
        raise ReadError("encoding_required") from None
    if "\x00" in text:
        raise ReadError("binary_text")
    if suffix == ".csv":
        try:
            dialect = csv.Sniffer().sniff(text[:65536], delimiters=";,\t|")
        except csv.Error:
            dialect = csv.excel
        rows = csv.reader(io.StringIO(text, newline=""), dialect, strict=True)
        result = []
        for csv_row in rows:
            result.append({"locator": f"row:{rows.line_num}", "values": csv_row})
            if len(result) > MAX_UNITS:
                raise ReadError("row_limit")
        return result, {"encoding": encoding}
    if suffix == ".json":
        strict_json(text)
    if suffix == ".xml":
        from defusedxml.ElementTree import fromstring

        try:
            fromstring(content, forbid_dtd=True)
        except Exception:
            raise ReadError("unsafe_xml") from None
    return [{"locator": f"lines:1-{len(text.splitlines())}", "text": text}], {"encoding": encoding}


def read_document(
    source: Path,
    output: Path,
    encoding: str = "utf-8-sig",
    *,
    pages: str | None = None,
    max_bytes: int = MAX_BYTES,
    expected_sha256: str | None = None,
) -> dict[str, Any]:
    content = checked_bytes(source, max_bytes, expected_sha256)
    output = private_directory(output)
    try:
        units, metadata = _read(content, source.suffix.lower(), output, encoding, pages)
        if len(units) > MAX_UNITS:
            raise ReadError("unit_limit")
        result = {
            "schema_version": "loki.local-reading/1",
            "status": "requires_harness_review",
            "untrusted_document_data": True,
            "source": {
                "filename": source.name,
                "sha256": hashlib.sha256(content).hexdigest(),
                "size_bytes": len(content),
            },
            "units": units,
            **metadata,
        }
        write_json(output / "manifest.json", result)
        return result
    except ReadError:
        raise
    except Exception:
        raise ReadError("document_read_failed") from None
